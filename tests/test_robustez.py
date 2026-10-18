from __future__ import annotations

import json
import os

os.environ.setdefault("OLLAMAPILOT_VAULT", "memory")

import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx

from ollamapilot import paths, undo
from ollamapilot.agent import Agent, Cancelled, estimate_tokens, fit_with_report, repair_transcript, request_index
from ollamapilot.attachments import describe, read_attachment
from ollamapilot.config import Settings
from ollamapilot.logbook import log_turn, redact
from ollamapilot.ollama_client import OllamaClient, OllamaError, describe_capabilities, parse_capabilities
from ollamapilot.tools.base import validate_arguments
from ollamapilot.tools.database import guard_select
from ollamapilot.tools.desktop import coerce_actions
from ollamapilot.tools.filesystem import DeletePath, EditFile, UndoChange, WriteFile
from ollamapilot.tools.programs import RunProgram


def call(name: str, **arguments) -> dict:
    return {"function": {"name": name, "arguments": arguments}}


def tools_step(*calls, content: str = "") -> dict:
    return {"content": content, "tool_calls": list(calls)}


def answer(text: str, **metrics) -> dict:
    return {"content": text, "tool_calls": [], "metrics": metrics}


class RecordingClient:
    """Cliente falso: devolve respostas prontas e guarda o que o agente enviou."""

    def __init__(self, steps: list, caps: dict | None = None) -> None:
        self.steps = list(steps)
        self.requests: list[dict] = []
        self.caps = caps

    def capabilities(self) -> dict:
        if self.caps is None:
            raise AttributeError("sem metadados")
        return self.caps

    def chat(self, messages, tools, on_event, cancel):
        self.requests.append({"messages": json.loads(json.dumps(messages)), "tools": tools})
        if not self.steps:
            return answer("fim")
        step = self.steps.pop(0)
        if isinstance(step, BaseException):
            raise step
        if callable(step):
            step = step(messages, tools)
        return step


def assert_valid(test: unittest.TestCase, transcript: list[dict]) -> None:
    """Toda chamada de ferramenta tem resultado, e nenhum resultado fica sem chamada."""
    index = 0
    while index < len(transcript):
        message = transcript[index]
        test.assertNotEqual(message.get("role"), "tool", f"resultado solto na posição {index}")
        calls = message.get("tool_calls") or []
        index += 1
        seen = 0
        while index < len(transcript) and transcript[index].get("role") == "tool":
            seen += 1
            index += 1
        test.assertGreaterEqual(seen, len(calls), f"chamada sem resultado antes da posição {index}")


def run(agent: Agent, text: str = "faça", approve=lambda _p: "once", cancel=lambda: False, attachments=None):
    events: list[tuple[str, object]] = []
    result = agent.turn(text, attachments or [], lambda kind, data: events.append((kind, data)), approve, cancel)
    return result, events


class ArgumentValidationTests(unittest.TestCase):
    def test_required_and_types(self) -> None:
        params = WriteFile(Settings()).parameters
        cleaned, errors = validate_arguments(params, {"caminho": r"C:\x\a.txt"})
        self.assertTrue(any("conteudo" in item for item in errors))
        cleaned, errors = validate_arguments(params, {"caminho": r"C:\x\a.txt", "conteudo": 12})
        self.assertEqual(errors, [])
        self.assertEqual(cleaned["conteudo"], "12")
        params = RunProgram(Settings(), lambda: False).parameters
        cleaned, errors = validate_arguments(params, {"programa": "git", "timeout_segundos": "30", "aguardar": "false"})
        self.assertEqual(errors, [])
        self.assertEqual(cleaned["timeout_segundos"], 30)
        self.assertIs(cleaned["aguardar"], False)
        _, errors = validate_arguments(params, {"programa": "git", "timeout_segundos": "muito"})
        self.assertTrue(errors)

    def test_write_without_content_is_refused_before_approval(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "vazio.txt"
            client = RecordingClient([tools_step(call("gravar_arquivo", caminho=str(target))), answer("Faltou o conteúdo.")])
            asked = []
            result, _ = run(Agent(Settings(model="t"), client), approve=lambda p: asked.append(p) or "once")
            self.assertEqual(result, "Faltou o conteúdo.")
            self.assertEqual(asked, [])
            self.assertFalse(target.exists())
            tool = client.requests[1]["messages"][-1]
            self.assertIn("conteudo", tool["content"])
            self.assertIn("orientacao", tool["content"])

    def test_malformed_json_arguments_are_reported(self) -> None:
        bad = {"function": {"name": "ler_arquivo", "arguments": "{caminho: sem aspas"}}
        client = RecordingClient([tools_step(bad), answer("Argumentos ruins.")])
        agent = Agent(Settings(model="t"), client)
        run(agent)
        tool = [m for m in agent.transcript if m["role"] == "tool"][0]
        self.assertIn("JSON válido", tool["content"])
        assert_valid(self, agent.transcript)


class LoopTests(unittest.TestCase):
    def test_reread_after_edit_is_not_a_duplicate(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "a.txt"
            target.write_text("valor = 1\n", encoding="utf-8")
            read = call("ler_arquivo", caminho=str(target))
            client = RecordingClient(
                [
                    tools_step(read),
                    tools_step(call("editar_arquivo", caminho=str(target), buscar="valor = 1", substituir="valor = 2")),
                    tools_step(json.loads(json.dumps(read))),
                    answer("Conferido."),
                ]
            )
            agent = Agent(Settings(model="t"), client)
            agent.workspace = Path(folder)
            run(agent)
            results = [m["content"] for m in agent.transcript if m["role"] == "tool"]
            self.assertIn("valor = 2", results[-1])
            self.assertNotIn("repetida", results[-1])

    def test_repeated_failure_stops_early(self) -> None:
        missing = call("ler_arquivo", caminho=r"C:\nao\existe\nada.txt")
        steps = [tools_step(json.loads(json.dumps(missing))) for _ in range(10)] + [answer("Não encontrei o arquivo.")]
        client = RecordingClient(steps)
        agent = Agent(Settings(model="t", max_tool_rounds=16), client)
        result, _ = run(agent)
        self.assertLessEqual(len(client.requests), 3)
        self.assertEqual(agent.last_trace["desfecho"], "sem_progresso")
        self.assertIn("exista", result)
        repeated = [m for m in agent.transcript if m["role"] == "tool" and "já falhou" in m["content"]]
        self.assertTrue(repeated)
        assert_valid(self, agent.transcript)

    def test_round_limit_forces_a_summary(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            steps = [tools_step(call("listar_pasta", caminho=folder + f"\\p{n}")) for n in range(3)]
            for n in range(3):
                (Path(folder) / f"p{n}").mkdir()
            client = RecordingClient(steps + [answer("Resumo do que foi feito.")])
            agent = Agent(Settings(model="t", max_tool_rounds=3), client)
            result, _ = run(agent)
            self.assertEqual(result, "Resumo do que foi feito.")
            self.assertEqual(agent.last_trace["desfecho"], "limite_rodadas")
            self.assertEqual(client.requests[-1]["tools"], [])
            self.assertIn("Limite de 3 rodadas", client.requests[-1]["messages"][-1]["content"])

    def test_too_many_calls_in_one_reply_all_get_results(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            calls = [call("listar_pasta", caminho=folder + "\\" + "x" * n) for n in range(1, 16)]
            client = RecordingClient([tools_step(*calls), answer("ok")])
            agent = Agent(Settings(model="t"), client)
            run(agent)
            tools = [m for m in agent.transcript if m["role"] == "tool"]
            self.assertEqual(len(tools), 15)
            self.assertIn("Limite de 12", tools[-1]["content"])
            assert_valid(self, agent.transcript)

    def test_server_failure_mid_turn_keeps_the_session_usable(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            client = RecordingClient(
                [tools_step(call("listar_pasta", caminho=folder)), OllamaError("A conexão com o Ollama caiu"), answer("De volta.")]
            )
            agent = Agent(Settings(model="t"), client)
            with self.assertRaises(OllamaError):
                run(agent)
            assert_valid(self, agent.transcript)
            self.assertEqual(agent.last_trace["desfecho"], "erro")
            result, _ = run(agent, "continue")
            self.assertEqual(result, "De volta.")
            assert_valid(self, agent.transcript)

    def test_crash_inside_approval_closes_every_call(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            first = call("gravar_arquivo", caminho=str(Path(folder) / "a.txt"), conteudo="a")
            second = call("gravar_arquivo", caminho=str(Path(folder) / "b.txt"), conteudo="b")
            client = RecordingClient([tools_step(first, second)])
            agent = Agent(Settings(model="t"), client)
            agent.workspace = Path(folder)

            def explode(_preview):
                raise RuntimeError("janela fechada")

            with self.assertRaises(RuntimeError):
                run(agent, approve=explode)
            self.assertEqual([m["role"] for m in agent.transcript], ["user", "assistant", "tool", "tool"])
            assert_valid(self, agent.transcript)

    def test_cancel_during_model_call_leaves_no_partial_answer(self) -> None:
        state = {"cancel": False}

        def slow(_messages, _tools):
            state["cancel"] = True
            return {"content": "meia resp", "tool_calls": [], "cancelled": True}

        client = RecordingClient([slow, answer("segunda")])
        agent = Agent(Settings(model="t"), client)
        with self.assertRaises(Cancelled):
            run(agent, cancel=lambda: state["cancel"])
        self.assertEqual([m["role"] for m in agent.transcript], ["user"])
        state["cancel"] = False
        result, _ = run(agent, "de novo", cancel=lambda: state["cancel"])
        self.assertEqual(result, "segunda")


class ContextTests(unittest.TestCase):
    def test_request_survives_a_long_tool_loop(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            files = []
            for n in range(6):
                path = Path(folder) / f"grande{n}.txt"
                path.write_text(("linha de dados %d\n" % n) * 900, encoding="utf-8")
                files.append(path)
            steps = [tools_step(call("ler_arquivo", caminho=str(path))) for path in files] + [answer("Pronto.")]
            client = RecordingClient(steps)
            settings = Settings(model="t", num_ctx=8192, max_tool_chars=24_000)
            agent = Agent(settings, client)
            agent.transcript = [
                {"role": "user", "content": r"Antes: use E:\obra\saida\relatorio.md " + "y" * 30000},
                {"role": "assistant", "content": "Combinado, vou usar esse caminho."},
            ]
            run(agent, "PEDIDO-ATUAL resuma os arquivos grandes")
            for request in client.requests:
                users = [m for m in request["messages"] if m["role"] == "user"]
                self.assertTrue(any("PEDIDO-ATUAL" in m["content"] for m in users))
                body = request["messages"][1:]
                self.assertLessEqual(estimate_tokens(body), 8192)
            later = client.requests[-1]["messages"]
            self.assertTrue(any("relatorio.md" in m["content"] for m in later))
            self.assertTrue(any("Resumo automático" in m["content"] for m in later))

    def test_nudge_does_not_strip_the_attachment(self) -> None:
        client = RecordingClient(
            [answer("Desculpe, mas não posso salvar arquivos."), answer("Certo, li o anexo.")]
        )
        agent = Agent(Settings(model="t"), client)
        attachment = {"name": "dados.csv", "text": "ANEXO-123;valor", "tipo": "texto", "codificacao": "utf-8"}
        run(agent, r"Crie o relatório e salve em C:\temp\saida", attachments=[attachment])
        second = client.requests[1]["messages"]
        self.assertTrue(any("ANEXO-123" in m["content"] for m in second))

    def test_old_images_are_not_resent(self) -> None:
        client = RecordingClient([answer("vi"), answer("ok")])
        agent = Agent(Settings(model="t"), client)
        image = {"name": "tela.png", "base64": "aGVsbG8=", "tipo": "imagem"}
        run(agent, "veja a imagem", attachments=[image])
        run(agent, "e agora?")
        self.assertTrue(any(m.get("images") for m in client.requests[0]["messages"]))
        self.assertFalse(any(m.get("images") for m in client.requests[1]["messages"]))
        self.assertTrue(any("tela.png" in m["content"] for m in client.requests[1]["messages"]))

    def test_summary_note_is_internal_and_keeps_paths(self) -> None:
        old = [
            {"role": "user", "content": r"Grave em D:\proj\app\main.py"},
            {"role": "assistant", "content": "", "tool_calls": [call("gravar_arquivo", caminho=r"D:\proj\app\main.py", conteudo="x")]},
            {"role": "tool", "tool_name": "gravar_arquivo", "content": '{"caminho": "D:\\\\proj\\\\app\\\\main.py"}'},
            {"role": "assistant", "content": "Gravei o main.py."},
            {"role": "user", "content": "Revise isto: " + "z" * 40000},
            {"role": "assistant", "content": "Revisado."},
            {"role": "user", "content": "agora teste"},
        ]
        fitted, report = fit_with_report(old, 4096)
        self.assertGreater(report["resumidas"], 0)
        self.assertEqual(request_index(fitted), len(fitted) - 1)
        note = fitted[0]
        self.assertTrue(note.get("interno"))
        self.assertIn(r"D:\proj\app\main.py", note["content"])
        self.assertIn("gravar_arquivo", note["content"])
        assert_valid(self, repair_transcript(fitted))


class CapabilityTests(unittest.TestCase):
    def test_parse_show_payload(self) -> None:
        caps = parse_capabilities(
            {"capabilities": ["completion", "tools", "thinking"], "model_info": {"qwen3.context_length": 40960}}
        )
        self.assertTrue(caps["ferramentas"])
        self.assertFalse(caps["visao"])
        self.assertEqual(caps["contexto"], 40960)
        self.assertIn("sem visão", describe_capabilities(caps))
        self.assertFalse(parse_capabilities({"modelfile": "x"})["conhecido"])

    def test_model_without_tools_gets_no_tools(self) -> None:
        client = RecordingClient([answer("Só texto.")], caps={"conhecido": True, "ferramentas": False, "visao": False})
        agent = Agent(Settings(model="t"), client)
        image = {"name": "foto.jpg", "base64": "aGVsbG8=", "tipo": "imagem"}
        result, events = run(agent, "descreva", attachments=[image])
        self.assertEqual(result, "Só texto.")
        self.assertEqual(client.requests[0]["tools"], [])
        self.assertFalse(any(m.get("images") for m in client.requests[0]["messages"]))
        self.assertIn("não aceita imagem", client.requests[0]["messages"][-1]["content"])
        self.assertTrue(any(kind == "trace" and data.get("tool") == "modelo" for kind, data in events))

    def test_remote_model_metadata_is_not_requested(self) -> None:
        client = OllamaClient(Settings(model="gpt-oss:120b", model_remote=True, model_source="online", cloud_api_key="abc"))
        with patch.object(OllamaClient, "_post", side_effect=AssertionError("não deveria consultar")):
            caps = client.capabilities()
        self.assertFalse(caps["conhecido"])

    def test_chat_payload_adapts_to_capabilities(self) -> None:
        sent = {}

        class Response:
            status_code = 200

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def iter_lines(self):
                yield json.dumps({"message": {"content": "oi"}, "done": True, "prompt_eval_count": 10})

            def close(self):
                pass

        class Client:
            def __init__(self, *args, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def stream(self, method, url, json=None, headers=None):
                sent.update(json)
                return Response()

        settings = Settings(model="m", think="high", num_ctx=32768)
        client = OllamaClient(settings)
        client._caps[(False, client._base(), "m")] = {"conhecido": True, "raciocinio": False, "contexto": 8192}
        with patch("ollamapilot.ollama_client.httpx.Client", Client):
            reply = client.chat([{"role": "user", "content": "oi"}], [], lambda *_: None, lambda: False)
        self.assertEqual(reply["content"], "oi")
        self.assertNotIn("think", sent)
        self.assertEqual(sent["options"]["num_ctx"], 8192)

    def test_cancel_interrupts_a_silent_stream(self) -> None:
        closed = threading.Event()

        class Response:
            status_code = 200

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def iter_lines(self):
                closed.wait(5)
                raise httpx.ReadError("fechado")

            def close(self):
                closed.set()

        class Client:
            def __init__(self, *args, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def stream(self, *args, **kwargs):
                return Response()

        start = time.perf_counter()
        deadline = start + 0.3
        with patch("ollamapilot.ollama_client.httpx.Client", Client):
            reply = OllamaClient(Settings(model="m")).chat([], [], lambda *_: None, lambda: time.perf_counter() > deadline)
        self.assertTrue(reply["cancelled"])
        self.assertLess(time.perf_counter() - start, 2)


class FileToolTests(unittest.TestCase):
    def test_undo_restores_and_removes(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "nota.txt"
            settings = Settings()
            WriteFile(settings).run({"caminho": str(target), "conteudo": "um"})
            WriteFile(settings).run({"caminho": str(target), "conteudo": "dois"})
            preview = UndoChange().preview({"caminho": str(target)})
            self.assertIn("versão anterior", preview["detail"])
            UndoChange().run({"caminho": str(target)})
            self.assertEqual(target.read_text(encoding="utf-8"), "um")
            UndoChange().run({"caminho": str(target)})
            self.assertFalse(target.exists())
            with self.assertRaises(ValueError):
                UndoChange().preview({"caminho": str(target)})

    def test_deleted_file_can_be_restored(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "apagar.txt"
            target.write_text("guardado", encoding="utf-8")
            preview = DeletePath().preview({"caminho": str(target)})
            self.assertIn("8 bytes", preview["detail"])
            DeletePath().run({"caminho": str(target)})
            self.assertFalse(target.exists())
            UndoChange().run({"caminho": str(target)})
            self.assertEqual(target.read_text(encoding="utf-8"), "guardado")

    def test_drive_root_cannot_be_deleted(self) -> None:
        with self.assertRaises(ValueError):
            DeletePath().preview({"caminho": "C:\\", "recursivo": True})

    def test_edit_shows_a_diff_and_handles_crlf(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "app.py"
            target.write_bytes(b"a = 1\r\nb = 2\r\nc = 3\r\n")
            edit = EditFile(Settings())
            arguments = {"caminho": str(target), "buscar": "a = 1\nb = 2", "substituir": "a = 1\nb = 20"}
            preview = edit.preview(arguments)
            self.assertIn("-b = 2", preview["detail"])
            self.assertIn("+b = 20", preview["detail"])
            edit.run(arguments)
            self.assertEqual(target.read_bytes(), b"a = 1\r\nb = 20\r\nc = 3\r\n")
            target.write_text("x\nx\n", encoding="utf-8")
            with self.assertRaises(ValueError) as caught:
                edit.preview({"caminho": str(target), "buscar": "x", "substituir": "y"})
            self.assertIn("2 vezes", str(caught.exception))

    def test_replace_preview_shows_what_changes(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "index.html"
            target.write_text("<h1>Velho</h1>\n", encoding="utf-8")
            preview = WriteFile(Settings()).preview({"caminho": str(target), "conteudo": "<h1>Novo</h1>\n"})
            self.assertEqual(preview["title"], "Substituir arquivo")
            self.assertIn("-<h1>Velho</h1>", preview["detail"])
            self.assertIn("+<h1>Novo</h1>", preview["detail"])

    def test_backups_stay_in_the_test_storage(self) -> None:
        self.assertTrue(str(paths.BACKUPS).startswith(os.environ["OLLAMAPILOT_STORAGE"]))


class ProgramAndPolicyTests(unittest.TestCase):
    def test_nonzero_exit_is_not_success(self) -> None:
        fake = {"exit_code": 2, "timed_out": False, "cancelled": False, "stdout": "", "stderr": "erro", "truncated": False}
        tool = RunProgram(Settings(), lambda: False)
        with patch("ollamapilot.tools.programs.run_process", return_value=dict(fake)), patch(
            "ollamapilot.tools.programs._program", return_value=r"C:\bin\x.exe"
        ):
            result = tool.run({"programa": "x"})
        self.assertFalse(result["sucesso"])
        self.assertIn("código 2", result["mensagem"])
        with patch("ollamapilot.tools.programs.run_process", return_value={**fake, "exit_code": -1, "timed_out": True}), patch(
            "ollamapilot.tools.programs._program", return_value=r"C:\bin\x.exe"
        ):
            result = tool.run({"programa": "x", "timeout_segundos": 5})
        self.assertIn("não terminou em 5 s", result["mensagem"])

    def test_select_guard_blocks_side_effects(self) -> None:
        for sql in ("SELECT pg_terminate_backend(123)", "SELECT * INTO copia FROM clientes", "select pg_sleep(100)"):
            with self.assertRaises(ValueError):
                guard_select(sql)
        self.assertTrue(guard_select("SELECT nome FROM clientes WHERE ativo"))

    def test_window_actions_are_checked_before_approval(self) -> None:
        with self.assertRaises(ValueError):
            coerce_actions({"acoes": [{"acao": "formatar"}]})
        with self.assertRaises(ValueError):
            coerce_actions({"acoes": [{"acao": "ler"}] * 30})


class AttachmentTests(unittest.TestCase):
    def test_office_and_pdf_text(self) -> None:
        from docx import Document
        from fpdf import FPDF
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            docx_path = Path(folder) / "contrato.docx"
            document = Document()
            document.add_paragraph("Cláusula primeira")
            document.save(str(docx_path))
            xlsx_path = Path(folder) / "vendas.xlsx"
            book = Workbook()
            book.active.append(["produto", "total"])
            book.active.append(["mouse", 30])
            book.save(str(xlsx_path))
            pdf_path = Path(folder) / "manual.pdf"
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.cell(0, 10, "Pagina com texto")
            pdf.add_page()
            pdf.cell(0, 10, "Segunda pagina")
            pdf.output(str(pdf_path))
            docx = read_attachment(docx_path)
            xlsx = read_attachment(xlsx_path)
            pdf_result = read_attachment(pdf_path)
        self.assertIn("Cláusula primeira", docx["anexo"]["text"])
        self.assertIn("mouse;30", xlsx["anexo"]["text"])
        self.assertIn("[página 2]", pdf_result["anexo"]["text"])
        self.assertEqual(pdf_result["anexo"]["paginas"], 2)
        self.assertIn("pdf", describe(pdf_result["anexo"]))

    def test_fake_image_and_binary_are_refused(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            fake = Path(folder) / "foto.png"
            fake.write_text("não sou imagem", encoding="utf-8")
            binary = Path(folder) / "dados.bin"
            binary.write_bytes(b"\x00\x01\x02" * 10)
            broken = Path(folder) / "quebrado.docx"
            broken.write_bytes(b"isto nao e zip")
            results = [read_attachment(fake), read_attachment(binary), read_attachment(broken)]
        self.assertTrue(all(not item["ok"] for item in results))
        self.assertIn("não é PNG", results[0]["erro"])
        self.assertIn("binário", results[1]["erro"])
        self.assertIn("corrompido", results[2]["erro"])

    def test_truncation_is_labelled(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "grande.txt"
            path.write_text("é" * 70_000, encoding="utf-8")
            item = read_attachment(path)["anexo"]
        self.assertTrue(item["truncado"])
        self.assertEqual(len(item["text"]), 50_000)
        self.assertIn("truncado em 50.000", describe(item))


class StorageAndLogTests(unittest.TestCase):
    def test_redact_common_secrets(self) -> None:
        text = redact("Authorization: Bearer abcdefghijkl123 senha=hunter2 chave sk-abcdefghijklmnop1234")
        self.assertNotIn("abcdefghijkl123", text)
        self.assertNotIn("hunter2", text)
        self.assertNotIn("sk-abcdefghijklmnop1234", text)

    def test_turn_trail_has_no_content(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            with patch("ollamapilot.logbook.LOGS", Path(folder)):
                client = RecordingClient([answer("RESPOSTA-SECRETA", prompt_eval_count=50, eval_count=7)])
                agent = Agent(Settings(model="m"), client)
                run(agent, "PEDIDO-PRIVADO")
                text = (Path(folder) / "turnos.jsonl").read_text(encoding="utf-8")
                record = json.loads(text.splitlines()[-1])
                agent.settings.trace_enabled = False
                client.steps.append(answer("outra"))
                run(agent, "mais")
                lines = (Path(folder) / "turnos.jsonl").read_text(encoding="utf-8").splitlines()
        self.assertNotIn("PEDIDO-PRIVADO", text)
        self.assertNotIn("RESPOSTA-SECRETA", text)
        self.assertEqual(record["desfecho"], "respondido")
        self.assertEqual(record["eval_count"], 7)
        self.assertEqual(record["prompt_eval_count"], 50)
        self.assertEqual(len(lines), 1)
        with tempfile.TemporaryDirectory() as folder:
            with patch("ollamapilot.logbook.LOGS", Path(folder)):
                log_turn({"modelo": "m", "cloud_api_key": "abc", "nota": "token=xyz123"})
                saved = (Path(folder) / "turnos.jsonl").read_text(encoding="utf-8")
        self.assertNotIn("abc", saved)
        self.assertNotIn("xyz123", saved)

    def test_settings_are_clamped_and_bad_encoding_is_parked(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "settings.json"
            path.write_text(json.dumps({"num_ctx": 10, "max_tool_rounds": 999, "temperature": 9}), encoding="utf-8")
            with patch("ollamapilot.config.SETTINGS_FILE", path):
                loaded = Settings.load()
            self.assertEqual(loaded.num_ctx, 2048)
            self.assertEqual(loaded.max_tool_rounds, 64)
            self.assertEqual(loaded.temperature, 2.0)
            path.write_bytes(b"\xff\xfe{\x00")
            with patch("ollamapilot.config.SETTINGS_FILE", path):
                fresh = Settings.load()
            self.assertEqual(fresh.model, "")
            self.assertTrue((Path(folder) / "settings.json.invalido").is_file())

    def test_second_invalid_copy_is_kept(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sessao.json"
            path.write_text("{", encoding="utf-8")
            paths.park_invalid(path)
            path.write_text("[", encoding="utf-8")
            paths.park_invalid(path)
            parked = sorted(item.name for item in Path(folder).iterdir())
        self.assertEqual(len(parked), 2)

    def test_atomic_write_keeps_the_old_copy_on_failure(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sessao.json"
            path.write_text('{"ok": 1}', encoding="utf-8")
            with patch("ollamapilot.paths.os.fsync", side_effect=OSError("disco cheio")):
                with self.assertRaises(OSError):
                    paths.write_text_atomic(path, '{"ok": 2}')
            self.assertEqual(path.read_text(encoding="utf-8"), '{"ok": 1}')
            self.assertEqual([item.name for item in Path(folder).iterdir()], ["sessao.json"])


class ModelSourceTests(unittest.TestCase):
    def setUp(self) -> None:
        self._folder = tempfile.TemporaryDirectory()
        self._path = Path(self._folder.name) / "settings.json"
        self._patch = patch("ollamapilot.config.SETTINGS_FILE", self._path)
        self._patch.start()

    def tearDown(self) -> None:
        self._patch.stop()
        self._folder.cleanup()

    def _load(self, raw: dict) -> Settings:
        self._path.write_text(json.dumps(raw), encoding="utf-8")
        return Settings.load()

    def test_old_settings_are_migrated(self) -> None:
        self.assertEqual(self._load({"cloud_api_key": "k", "model_remote": True}).model_source, "online")
        self.assertEqual(self._load({"cloud_api_key": "k", "model_remote": False}).model_source, "local")
        self.assertEqual(self._load({"model_source": "nuvem-toda"}).model_source, "local")
        self.assertEqual(self._load({"model_source": "online"}).model_source, "online")

    def test_local_mode_never_queries_ollama_com(self) -> None:
        client = OllamaClient(Settings(model_source="local", cloud_api_key="k"))
        with patch.object(client, "version", return_value={}), patch.object(client, "tags", return_value=[{"name": "a"}]), patch.object(
            client, "cloud_tags", side_effect=AssertionError("consultou a nuvem")
        ):
            payload = client.available_models()
        self.assertEqual(payload["remote"], [])
        online = OllamaClient(Settings(model_source="online", cloud_api_key="k"))
        with patch.object(online, "version", return_value={}), patch.object(online, "tags", return_value=[]), patch.object(
            online, "cloud_tags", return_value=[{"name": "b"}]
        ):
            self.assertEqual(online.available_models()["remote"], [{"name": "b"}])
        without_key = OllamaClient(Settings(model_source="online", cloud_api_key=""))
        with patch.object(without_key, "version", return_value={}), patch.object(without_key, "tags", return_value=[]), patch.object(
            without_key, "cloud_tags", side_effect=AssertionError("consultou sem chave")
        ):
            self.assertEqual(without_key.available_models()["remote"], [])

    def test_remote_chat_is_refused_in_local_mode(self) -> None:
        client = OllamaClient(Settings(model="gpt-oss:120b", model_remote=True, model_source="local", cloud_api_key="k"))
        with self.assertRaises(OllamaError) as caught:
            client.chat_target()
        self.assertIn("desligados", str(caught.exception))


class WindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        os.environ["OLLAMAPILOT_SELFTEST"] = "1"
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def _window(self, **options):
        from ollamapilot.ui.main_window import MainWindow

        window = MainWindow(Settings(model="sumiu", **options))
        window.client.capabilities = lambda *args, **kwargs: {"conhecido": False}
        return window

    def _settle(self, window) -> None:
        for job in list(window._jobs):
            job.wait(3000)
        self.app.processEvents()

    def test_cloud_model_is_never_picked_by_itself(self) -> None:
        window = self._window(model_source="online", cloud_api_key="chave")
        with patch.object(Settings, "save"):
            window._apply_connection({"models": [], "remote": [{"name": "gpt-oss:120b"}]})
            self.assertFalse(window.settings.model_remote)
            self.assertEqual(window.settings.model, "")
            window._apply_connection({"models": [{"name": "qwen3:8b"}], "remote": [{"name": "gpt-oss:120b"}]})
            self.assertEqual(window.settings.model, "qwen3:8b")
            self.assertFalse(window.settings.model_remote)
        self._settle(window)
        window.close()

    def test_local_only_mode_hides_online_models(self) -> None:
        window = self._window(model_source="local", cloud_api_key="chave")
        with patch.object(Settings, "save"):
            window._apply_connection({"models": [{"name": "qwen3:8b"}], "remote": [{"name": "gpt-oss:120b"}]})
        names = [window.models.itemText(i) for i in range(window.models.count())]
        self.assertEqual(names, ["qwen3:8b"])
        self.assertFalse(window.balance_button.isVisibleTo(window))
        self._settle(window)
        window.close()

    def test_online_mode_requires_the_api_key(self) -> None:
        from ollamapilot.ui.dialogs import SettingsDialog

        window = self._window(model_source="local")
        dialog = SettingsDialog(window.settings, window.client, window)
        dialog.source.setCurrentIndex(dialog.source.findData("online"))
        dialog.cloud_key.setText("")
        with patch.object(Settings, "save") as save:
            dialog._save()
        save.assert_not_called()
        self.assertEqual(window.settings.model_source, "local")
        self.assertIn("API Key", dialog.feedback.text())
        dialog.cloud_key.setText("chave")
        with patch.object(Settings, "save"), patch.object(window.client, "version", return_value={"version": "x"}):
            dialog._save()
        self.assertEqual(window.settings.model_source, "online")
        self.assertTrue(window.settings.online_models())
        dialog.close()
        self._settle(window)
        window.close()

    def test_garbage_session_is_sanitized(self) -> None:
        paths.SESSION_FILE.parent.mkdir(parents=True, exist_ok=True)
        paths.SESSION_FILE.write_text(
            json.dumps(
                {
                    "transcript": [
                        "lixo",
                        {"role": "assistant", "content": "", "tool_calls": [call("ler_arquivo", caminho="C:\\a")]},
                    ],
                    "display": ["lixo", {"role": "hacker", "content": 5, "trace": "x", "metrics": [1]}],
                    "autorizacoes": ["arquivos", 3],
                }
            ),
            encoding="utf-8",
        )
        try:
            window = self._window()
            self.assertEqual(len(window.display), 1)
            self.assertEqual(window.display[0]["role"], "assistant")
            self.assertEqual(window.agent.chat_grants, {"arquivos"})
            assert_valid(self, window.agent.transcript)
            self._settle(window)
            window.close()
        finally:
            paths.SESSION_FILE.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
