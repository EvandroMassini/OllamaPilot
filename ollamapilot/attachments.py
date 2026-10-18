from __future__ import annotations

import base64
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
DOCUMENTS = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".odt", ".epub", ".rtf", ".xls"}
AUDIO = {".wav", ".mp3", ".m4a", ".flac"}
MAX_BYTES = 10 * 1024 * 1024
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_TEXT = 50_000
MAX_INDEX_TEXT = 500_000
MAX_PAGES = 200
MAX_ROWS = 2000
PDF_IMAGE_PAGES = 4
PDF_IMAGE_EDGE = 1400

_IMAGE_MAGIC = (b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"GIF87a", b"GIF89a")


def read_attachment(path: Path, limit: int = MAX_TEXT) -> dict:
    """Lê um anexo. A falha de um arquivo não impede a leitura dos outros."""
    try:
        size = path.stat().st_size
    except OSError as error:
        return {"ok": False, "nome": path.name, "erro": _access_error(path, error)}
    if not path.is_file():
        return {"ok": False, "nome": path.name, "erro": f"{path.name} não é um arquivo."}
    if size > MAX_BYTES:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} passa de 10 MB e não foi anexado."}
    suffix = path.suffix.lower()
    try:
        if suffix in IMAGES:
            return _image(path, size)
        if suffix in AUDIO:
            return _audio_file(path, size)
        if suffix in DOCUMENTS:
            return _document(path, size, suffix, limit)
        with path.open("rb") as handle:
            data = handle.read(limit * 4 + 1)
    except OSError as error:
        return {"ok": False, "nome": path.name, "erro": _access_error(path, error)}
    if b"\x00" in data[:8192]:
        return {
            "ok": False,
            "nome": path.name,
            "erro": f"{path.name} é binário ({suffix or 'sem extensão'}) e não foi anexado. Para um executável, peça ler_binario com o caminho.",
        }
    from ollamapilot.tools.filesystem import decode_text

    text, encoding = decode_text(data)
    truncated = size > len(data)
    return _text_result(path, text, "texto", size, truncated, limit, codificacao=encoding)


def extract_document_text(path: Path) -> str:
    """Texto de um documento comum, para a biblioteca. Levanta ValueError se não houver texto legível."""
    result = read_attachment(path, MAX_INDEX_TEXT)
    if not result["ok"]:
        raise ValueError(result["erro"])
    text = result["anexo"].get("text")
    if not text:
        raise ValueError(f"{path.name} não tem texto para indexar.")
    return text


def describe(item: dict) -> str:
    """Rótulo curto com tipo, codificação, páginas e truncamento."""
    parts = [str(item.get("tipo") or "texto")]
    if item.get("codificacao"):
        parts.append(str(item["codificacao"]))
    if item.get("paginas"):
        parts.append(f"{item['paginas']} páginas")
    if item.get("truncado"):
        parts.append(f"truncado em {len(str(item.get('text') or '')):,} caracteres".replace(",", "."))
    return ", ".join(parts)


def _image(path: Path, size: int) -> dict:
    if size > MAX_IMAGE_BYTES:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} passa de 8 MB; imagens maiores não são enviadas ao modelo."}
    data = path.read_bytes()
    valid = data.startswith(_IMAGE_MAGIC) or (data[:4] == b"RIFF" and data[8:12] == b"WEBP")
    if not valid:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} tem extensão de imagem, mas o conteúdo não é PNG, JPEG, GIF ou WebP."}
    return {
        "ok": True,
        "anexo": {
            "name": path.name,
            "base64": base64.b64encode(data).decode("ascii"),
            "tipo": "imagem",
            "bytes": size,
        },
    }


def _document(path: Path, size: int, suffix: str, limit: int) -> dict:
    if suffix == ".rtf":
        try:
            return _text_result(path, _rtf_text(path), "rtf", size, False, limit)
        except OSError:
            raise
        except Exception as error:
            return {"ok": False, "nome": path.name, "erro": f"{path.name} não pôde ser lido como rtf: {error}"}
    if suffix == ".xls":
        try:
            text, truncated = _xls_text(path, limit)
        except ImportError:
            return {
                "ok": False,
                "nome": path.name,
                "erro": "Para ler planilhas .xls, instale a biblioteca xlrd. O arquivo não foi lido.",
            }
        except Exception as error:
            return {"ok": False, "nome": path.name, "erro": f"{path.name} não pôde ser lido como xls: {error}"}
        return _text_result(path, text, "xls", size, truncated, limit)
    if suffix != ".pdf" and not zipfile.is_zipfile(path):
        return {"ok": False, "nome": path.name, "erro": f"{path.name} está corrompido ou não é um {suffix} válido."}
    try:
        if suffix == ".pdf":
            text, pages, truncated = _pdf_text(path, limit)
            if not text.replace("[página", "").strip(" \n0123456789]"):
                return {
                    "ok": False,
                    "nome": path.name,
                    "erro": f"{path.name} não tem texto extraível (pode ser digitalizado como imagem).",
                }
            return _text_result(path, text, "pdf", size, truncated, limit, paginas=pages)
        if suffix == ".docx":
            return _text_result(path, _docx_text(path), "docx", size, False, limit)
        if suffix == ".pptx":
            return _text_result(path, _pptx_text(path), "pptx", size, False, limit)
        if suffix == ".odt":
            return _text_result(path, _xml_text(path, "content.xml"), "odt", size, False, limit)
        if suffix == ".epub":
            return _text_result(path, _epub_text(path), "epub", size, False, limit)
        text, truncated = _xlsx_text(path, limit)
        return _text_result(path, text, "xlsx", size, truncated, limit)
    except ImportError as error:
        if suffix == ".xls":
            return {
                "ok": False,
                "nome": path.name,
                "erro": "Para ler planilhas .xls, instale a biblioteca xlrd. O arquivo não foi lido.",
            }
        return {"ok": False, "nome": path.name, "erro": f"Falta a biblioteca para ler {suffix}: {error.name}."}
    except (zipfile.BadZipFile, KeyError) as error:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} está corrompido ou não é um {suffix} válido ({error})."}
    except OSError:
        raise
    except Exception as error:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} não pôde ser lido como {suffix}: {error}"}


def _text_result(path: Path, text: str, kind: str, size: int, truncated: bool, limit: int, **extra) -> dict:
    if len(text) > limit:
        text = text[:limit]
        truncated = True
    anexo = {"name": path.name, "text": text, "tipo": kind, "truncado": truncated, "bytes": size}
    anexo.update({key: value for key, value in extra.items() if value})
    return {"ok": True, "anexo": anexo}


def pdf_page_images(path: Path, first: int, last: int) -> dict:
    """Imagens das páginas de um PDF sem texto. Até PDF_IMAGE_PAGES por chamada."""
    import base64
    import io

    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages)
    if first > total:
        raise ValueError(f"O PDF tem {total} páginas; a página {first} não existe.")
    last = min(last, total, first + PDF_IMAGE_PAGES - 1)
    found = []
    missing = []
    for number in range(first, last + 1):
        rasterized = False
        image = _largest_page_image(reader.pages[number - 1])
        if image is None:
            try:
                image = _raster_page(path, number)
                rasterized = image is not None
            except _RasterMissing as error:
                return {"paginas": total, "imagens": found, "sem_imagem": missing + [number], "dependencia": str(error)}
            except Exception:
                missing.append(number)
                continue
        if image is None:
            missing.append(number)
            continue
        image = image.convert("RGB")
        image.thumbnail((PDF_IMAGE_EDGE, PDF_IMAGE_EDGE))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        found.append(
            {
                "pagina": number,
                "largura": image.width,
                "altura": image.height,
                "base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
                "raster": rasterized,
            }
        )
    return {"paginas": total, "imagens": found, "sem_imagem": missing}


def _largest_page_image(page):
    try:
        images = list(page.images)
    except Exception:
        return None
    best = None
    best_area = 0
    for item in images:
        try:
            image = item.image
        except Exception:
            continue
        area = int(image.width) * int(image.height)
        if area > best_area:
            best = image
            best_area = area
    return best


def _pdf_text(path: Path, limit: int) -> tuple[str, int, bool]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as error:
            raise ValueError("o PDF está protegido por senha") from error
    total = len(reader.pages)
    pieces: list[str] = []
    used = 0
    truncated = total > MAX_PAGES
    for number, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        text = (page.extract_text() or "").strip()
        block = f"[página {number}]\n{text}\n"
        pieces.append(block)
        used += len(block)
        if used > limit:
            truncated = True
            break
    return "\n".join(pieces), total, truncated


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for index, table in enumerate(document.tables, start=1):
        lines.append(f"[tabela {index}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx_text(path: Path, limit: int) -> tuple[str, bool]:
    from openpyxl import load_workbook

    formulas = load_workbook(str(path), data_only=False)
    cached = load_workbook(str(path), data_only=True)
    lines: list[str] = []
    used = 0
    truncated = False
    try:
        for sheet_f, sheet_v in zip(formulas.worksheets, cached.worksheets):
            lines.append(f"[planilha {sheet_f.title}]")
            last_row = min(sheet_f.max_row or 1, MAX_ROWS)
            last_col = sheet_f.max_column or 1
            if (sheet_f.max_row or 0) > MAX_ROWS:
                truncated = True
            for row_index in range(1, last_row + 1):
                cells = []
                empty = True
                for col_index in range(1, last_col + 1):
                    shown = _xlsx_cell(sheet_f.cell(row_index, col_index).value, sheet_v.cell(row_index, col_index).value)
                    if shown:
                        empty = False
                    cells.append(shown)
                if empty:
                    continue
                line = ";".join(cells)
                lines.append(line)
                used += len(line) + 1
                if used > limit:
                    truncated = True
                    break
            if (sheet_f.max_row or 0) > MAX_ROWS:
                lines.append(f"… linhas além de {MAX_ROWS} omitidas")
            if used > limit:
                break
    finally:
        formulas.close()
        cached.close()
    return "\n".join(lines), truncated


def _xlsx_cell(formula, cached) -> str:
    if isinstance(formula, str) and formula.startswith("="):
        if cached is None or cached == "":
            return formula
        return f"{formula} ({cached})"
    shown = formula if formula is not None else cached
    if shown is None:
        return ""
    return str(shown)


class _RasterMissing(RuntimeError):
    pass


def _raster_page(path: Path, number: int):
    try:
        import pypdfium2 as pdfium
    except ImportError as error:
        raise _RasterMissing(
            "Falta a biblioteca pypdfium2 para rasterizar a página. O conteúdo não foi lido."
        ) from error
    pdf = pdfium.PdfDocument(str(path))
    try:
        page = pdf[number - 1]
        width = float(page.get_width() or 1)
        scale = min(2.0, PDF_IMAGE_EDGE / max(width, 1))
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil().convert("RGB")
        image.thumbnail((PDF_IMAGE_EDGE, PDF_IMAGE_EDGE))
        return image
    finally:
        pdf.close()


def ocr_image(image) -> str | None:
    """Texto do Tesseract, se estiver no PATH. None quando não há OCR."""
    import shutil
    import subprocess
    import tempfile

    if not shutil.which("tesseract"):
        return None
    with tempfile.TemporaryDirectory() as folder:
        png = Path(folder) / "pagina.png"
        image.save(png, format="PNG")
        for language in ("por+eng", "eng"):
            try:
                completed = subprocess.run(
                    ["tesseract", str(png), "stdout", "-l", language],
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=60,
                    check=False,
                )
            except OSError:
                return None
            text = (completed.stdout or "").strip()
            if completed.returncode == 0 and text:
                return text
    return None


def _audio_file(path: Path, size: int) -> dict:
    if size > MAX_BYTES:
        return {"ok": False, "nome": path.name, "erro": f"{path.name} passa de 10 MB e não foi anexado."}
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return {
        "ok": True,
        "anexo": {
            "name": path.name,
            "audio": encoded,
            "tipo": "audio",
            "formato": path.suffix.lower().lstrip("."),
            "bytes": size,
        },
    }


def _local_name(tag: str) -> str:
    if "}" in tag:
        return tag.rsplit("}", 1)[-1]
    return tag


def _xml_texts(data: bytes) -> list[str]:
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError:
        return []
    found = []
    for node in root.iter():
        if _local_name(node.tag) in {"t", "p", "h", "span"} and node.text and node.text.strip():
            found.append(node.text.strip())
    return found


def _pptx_text(path: Path) -> str:
    lines = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
        for name in slides:
            number = re.search(r"(\d+)", Path(name).stem)
            label = number.group(1) if number else "?"
            body = _xml_texts(archive.read(name))
            notes_name = f"ppt/notesSlides/notesSlide{label}.xml"
            notes = _xml_texts(archive.read(notes_name)) if notes_name in archive.namelist() else []
            lines.append(f"[slide {label}]")
            lines.extend(body)
            if notes:
                lines.append(f"[notas {label}]")
                lines.extend(notes)
    return "\n".join(lines)


def _xml_text(path: Path, member: str) -> str:
    with zipfile.ZipFile(path) as archive:
        if member not in archive.namelist():
            raise ValueError(f"não encontrei {member}")
        return "\n".join(_xml_texts(archive.read(member)))


def _epub_text(path: Path) -> str:
    pieces = []
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if name.lower().endswith((".xhtml", ".html", ".htm")) and not name.startswith("META-INF/")
        ]
        for name in names:
            raw = archive.read(name).decode("utf-8", errors="replace")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = re.sub(r"\s+", " ", text).strip()
            if text:
                pieces.append(text)
    return "\n".join(pieces)


def _rtf_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    text = re.sub(r"\\'[0-9a-fA-F]{2}", lambda match: bytes.fromhex(match.group(0)[2:]).decode("cp1252", "replace"), text)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("\\{", "{").replace("\\}", "}").replace("\\\\", "\\")
    text = re.sub(r"[{}]", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _xls_text(path: Path, limit: int) -> tuple[str, bool]:
    try:
        import xlrd
    except ImportError:
        raise
    book = xlrd.open_workbook(str(path))
    lines: list[str] = []
    used = 0
    truncated = False
    for sheet in book.sheets():
        lines.append(f"[planilha {sheet.name}]")
        last = min(sheet.nrows, MAX_ROWS)
        if sheet.nrows > MAX_ROWS:
            truncated = True
        for row_index in range(last):
            cells = []
            empty = True
            for col_index in range(sheet.ncols):
                value = sheet.cell_value(row_index, col_index)
                shown = "" if value in (None, "") else str(value)
                if shown:
                    empty = False
                cells.append(shown)
            if empty:
                continue
            line = ";".join(cells)
            lines.append(line)
            used += len(line) + 1
            if used > limit:
                truncated = True
                break
        if sheet.nrows > MAX_ROWS:
            lines.append(f"… linhas além de {MAX_ROWS} omitidas")
        if used > limit:
            break
    return "\n".join(lines), truncated


def _access_error(path: Path, error: OSError) -> str:
    reason = error.strerror or "arquivo inacessível"
    return f"{path.name} não pôde ser lido ({reason})."
