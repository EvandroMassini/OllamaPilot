from __future__ import annotations

import csv
import io
import json
import time
import uuid
import zipfile
from pathlib import Path

from ollamapilot.config import Settings
from ollamapilot.paths import ARTIFACTS
from ollamapilot.tools.base import Tool, schema, text_prop, tool

_EXTENSIONS = {
    "txt",
    "md",
    "csv",
    "json",
    "xml",
    "html",
    "css",
    "js",
    "ts",
    "php",
    "py",
    "sql",
    "java",
    "c",
    "cpp",
    "h",
    "yaml",
    "yml",
    "pdf",
    "docx",
    "xlsx",
    "zip",
}


@tool
class CreateArtifact(Tool):
    name = "criar_arquivo"
    group = "biblioteca"
    order = 210
    scope_label = "gerar arquivos na biblioteca interna"
    description = (
        "Cria um documento para download em Arquivos gerados (texto, código, PDF, DOCX, XLSX via CSV, ZIP via arquivos), "
        "só quando o usuário não deu caminho; com caminho, use gravar_arquivo. Some em 24 h. Retorno: {caminho, nome}."
    )
    sensitive = True
    grant_key = "biblioteca"
    parameters = schema(
        {
            "nome": text_prop("Nome do arquivo com extensão, como relatorio.pdf."),
            "conteudo": text_prop("Conteúdo do arquivo. Em XLSX, envie CSV."),
            "arquivos": {
                "type": "array",
                "description": "Para ZIP: lista de objetos com nome e conteudo.",
                "items": {"type": "object"},
            },
        },
        ["nome"],
    )

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def preview(self, arguments: dict) -> dict:
        name = _safe_name(str(arguments.get("nome") or ""))
        return {
            "title": "Gerar arquivo",
            "summary": name,
            "detail": f"O arquivo {name} será criado na biblioteca do OllamaPilot.",
        }

    def run(self, arguments: dict) -> dict:
        name = _safe_name(str(arguments.get("nome") or ""))
        content = str(arguments.get("conteudo") or "")
        if len(content.encode("utf-8")) > self.settings.max_write_bytes:
            raise ValueError("O conteúdo excede o limite.")
        folder = ARTIFACTS / uuid.uuid4().hex
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / name
        extension = target.suffix.lower().lstrip(".")
        if extension == "pdf":
            _pdf(target, content)
        elif extension == "docx":
            _docx(target, content)
        elif extension == "xlsx":
            _xlsx(target, content)
        elif extension == "zip":
            _zip(target, arguments.get("arquivos"), content)
        else:
            target.write_text(content, encoding="utf-8", newline="\n")
        meta = {"nome": name, "caminho": str(target), "criado": time.time()}
        (folder / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        return {"caminho": str(target), "nome": name, "mensagem": "Arquivo disponível em Arquivos gerados."}


def list_artifacts() -> list[dict]:
    purge()
    rows = []
    if not ARTIFACTS.is_dir():
        return rows
    for meta_path in ARTIFACTS.glob("*/meta.json"):
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        path = Path(str(meta.get("caminho") or ""))
        if path.is_file():
            rows.append(meta)
    rows.sort(key=lambda item: item.get("criado") or 0, reverse=True)
    return rows


def purge(ttl_seconds: int = 24 * 60 * 60) -> None:
    if not ARTIFACTS.is_dir():
        return
    limit = time.time() - ttl_seconds
    for meta_path in ARTIFACTS.glob("*/meta.json"):
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            created = float(meta.get("criado") or 0)
        except (OSError, json.JSONDecodeError, TypeError, ValueError):
            continue
        if created < limit:
            import shutil

            shutil.rmtree(meta_path.parent, ignore_errors=True)


def _safe_name(name: str) -> str:
    name = Path(name).name.strip()
    if not name or name in {".", ".."} or any(char in name for char in '<>:"|?*'):
        raise ValueError("Nome de arquivo inválido.")
    extension = Path(name).suffix.lower().lstrip(".")
    if extension not in _EXTENSIONS:
        raise ValueError("Extensão não suportada na biblioteca. Para outro tipo, use gravar_arquivo.")
    return name


def _pdf(target: Path, content: str) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font = Path(r"C:\Windows\Fonts\arial.ttf")
    if font.is_file():
        pdf.add_font("Body", "", str(font))
        pdf.set_font("Body", size=11)
    else:
        pdf.set_font("Helvetica", size=11)
        content = content.encode("latin-1", errors="replace").decode("latin-1")
    pdf.multi_cell(0, 6, content)
    pdf.output(str(target))


def _docx(target: Path, content: str) -> None:
    from docx import Document

    document = Document()
    for paragraph in content.splitlines() or [""]:
        document.add_paragraph(paragraph)
    document.save(str(target))


def _xlsx(target: Path, content: str) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    for row in csv.reader(io.StringIO(content)):
        sheet.append(row)
    book.save(str(target))


def _zip(target: Path, files, content: str) -> None:
    items = files
    if not items:
        try:
            items = json.loads(content or "[]")
        except json.JSONDecodeError as error:
            raise ValueError("Para ZIP, envie a lista em arquivos, com nome e conteudo.") from error
    if not isinstance(items, list) or not items:
        raise ValueError("O ZIP precisa de ao menos um arquivo.")
    if len(items) > 30:
        raise ValueError("O ZIP aceita até 30 arquivos.")
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for item in items:
            if not isinstance(item, dict):
                raise ValueError("Cada item do ZIP precisa ter nome e conteudo.")
            name = Path(str(item.get("nome") or "")).name
            if not name:
                raise ValueError("Um item do ZIP está sem nome.")
            archive.writestr(name, str(item.get("conteudo") or ""))
